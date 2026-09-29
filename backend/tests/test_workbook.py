from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.pricing.pricecharting import normalize_pricecharting_url
from app.services.importer import detect_day_first, parse_bool


def _workbook() -> bytes:
    """A tracker in the style of many collectors' sheets: totals above a row-4 header,
    one tab per platform, PriceCharting links in HYPERLINK() formulas and cell links."""
    wb = Workbook()
    wb.active.title = "Totals"
    wb["Totals"]["A1"] = "Grand total"
    for name, games in {
        "NES": [
            ("1942", "Loose", 10, 0, 0, 1, "Game Store", "26/07/2016", 10, "formula", "nes/1942"),
            ("Airwolf", "Box", 10, 0, 10, 1, "Ebay", datetime(2016, 7, 11), 3, "legacy", "nes/airwolf"),
            ("Zelda", "SIB", 10, 10, 10, 2, "Garage Sale", "30/07/2016", 99.5, "cell", "nes/legend-of-zelda"),
        ],
        "N64": [("GoldenEye 007", "CIB", 10, 10, 10, 1, "Thrift Store", "01/02/2020", 25, None, None)],
    }.items():
        ws = wb.create_sheet(name)
        ws["C2"], ws["I2"], ws["J2"] = "Loose", "Paid", 123  # summary junk above the header
        header = [
            "",
            "Name",
            "Rating",
            "Cart",
            "Manual",
            "Box",
            "Amount",
            "Purchased",
            "Notes",
            "Date Added",
            "Paid",
        ]
        for c, h in enumerate(header, start=1):
            ws.cell(4, c, h or None)
        for r, (title, rating, cart, man, box, qty, src, added, paid, link, slug) in enumerate(
            games, start=5
        ):
            for c, v in enumerate(
                [None, title, rating, cart, man, box, qty, src, None, added, paid], start=1
            ):
                ws.cell(r, c, v)
            if link == "formula":
                ws.cell(r, 9, f'=HYPERLINK("https://www.pricecharting.com/game/{slug}","Lookup")')
            elif link == "legacy":
                ws.cell(r, 9, "Lookup").hyperlink = f"http://videogames.pricecharting.com/game/{slug}?q=x"
            elif link == "cell":
                ws.cell(r, 9, "Lookup").hyperlink = f"https://www.pricecharting.com/game/{slug}"
        ws.cell(40, 1, name)  # stray prefilled cell far below: not a row of data
    wb.create_sheet("Empty").sheet_state = "hidden"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_normalize_pricecharting_url():
    n = normalize_pricecharting_url
    assert n("http://videogames.pricecharting.com/game/nes/airwolf?q=airwolf") == (
        "https://www.pricecharting.com/game/nes/airwolf"
    )
    assert n("https://pricecharting.com/game/nes/1942#x") == "https://www.pricecharting.com/game/nes/1942"
    assert n("https://www.ebay.com/game/nes/1942") is None
    assert n("https://www.pricecharting.com/console/nes") is None
    assert n("javascript:alert(1)") is None


def test_helpers():
    assert parse_bool("10") is True and parse_bool("0") is False and parse_bool("7") is True
    assert detect_day_first([{"d": "26/07/2016"}, {"d": "01/02/2020"}], ["d"]) is True
    assert detect_day_first([{"d": "07/26/2016"}], ["d"]) is False
    assert detect_day_first([{"d": "2016-07-26"}], ["d"]) is False


def test_workbook_upload_and_import(admin):
    r = admin.post(
        "/api/import/workbook",
        files={
            "file": (
                "tracker.xlsx",
                _workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    sheets = {s["name"]: s for s in r.json()["sheets"]}
    assert sheets["Empty"]["hidden"] and sheets["Empty"]["rows"] == []
    nes = sheets["NES"]
    assert nes["header_row"] == 4
    assert nes["headers"] == [
        "Name", "Rating", "Cart", "Manual", "Box", "Amount", "Purchased", "Notes", "Date Added", "Paid",
        "Notes (link)", "Sheet name",
    ]  # fmt: skip
    assert len(nes["rows"]) == 3  # the stray cell on row 40 is in an unnamed column
    assert nes["rows"][0]["Name"] == "1942" and nes["rows"][0]["Sheet name"] == "NES"
    assert nes["rows"][1]["Date Added"] == "2016-07-11"  # real dates become ISO text

    rows = nes["rows"] + sheets["N64"]["rows"]
    mapping = admin.post("/api/import/suggest", json={"headers": nes["headers"], "rows": rows}).json()
    assert mapping == {
        "title": "Name",
        "platform": "Sheet name",
        "condition": "Rating",
        "quantity": "Amount",
        "has_item": "Cart",
        "has_manual": "Manual",
        "has_box": "Box",
        "acquired_from": "Purchased",
        "purchase_date": "Date Added",
        "purchase_price": "Paid",
        "pricecharting_url": "Notes (link)",
    }  # "Notes" is only "Lookup" link text, so it isn't mapped

    result = admin.post("/api/import/commit", json={"rows": rows, "mapping": mapping}).json()
    assert result["created_items"] == 4 and result["failed"] == 0
    items = {
        i["product"]["title"]: i
        for i in admin.get("/api/collection", params={"status": "owned"}).json()["items"]
    }
    assert items["1942"]["purchase_date"] == "2016-07-26"  # day-first detected
    airwolf = items["Airwolf"]
    assert (
        airwolf["condition"] == "loose" and airwolf["has_box"] and not airwolf["has_manual"]
    )  # "Box" rating
    assert airwolf["product"]["pricecharting_url"] == "https://www.pricecharting.com/game/nes/airwolf"
    assert airwolf["acquired_from"] == "Ebay"
    zelda = items["Zelda"]
    assert zelda["condition"] == "new" and zelda["quantity"] == 2 and zelda["purchase_price"] == "99.50"
    assert items["GoldenEye 007"]["product"]["platform"]["name"] == "Nintendo 64"
    assert admin.get("/api/collection/facets").json()["sources"] == [
        "Ebay",
        "Game Store",
        "Garage Sale",
        "Thrift Store",
    ]
    assert admin.get("/api/collection", params={"acquired_from": "Ebay"}).json()["total"] == 1


def test_workbook_rejects_non_excel(admin):
    r = admin.post(
        "/api/import/workbook", files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")}
    )
    assert r.status_code == 400


def test_link_matches_title():
    from app.services.importer import link_matches_title

    url = "https://www.pricecharting.com/game/{}".format
    ok = [
        ("Madden 98", "sega-genesis/madden-nfl-98"),
        ("SoulBlazer", "super-nintendo/soul-blazer"),
        ("Sim City", "super-nintendo/simcity"),
        ("Spider-man", "atari-2600/spiderman"),
        ("Mash", "atari-2600/m*a*s*h"),
        ("Madden 95", "sega-genesis/madden-nfl-%2795"),
        ("Baseball Simulator 1,000", "nes/baseball-simulator-1000"),
        ("Tiger Woods PGA Tour 07", "xbox/tiger-woods-2007"),
        ("Final Fantasy II", "super-nintendo/final-fantasy-2"),
        ("The Legend of Zelda - A Link to the Past", "super-nintendo/zelda-link-to-the-past"),
    ]
    wrong = [
        ("Stealth", "nes/rocket-ranger"),
        ("Turok 2 Seeds of Evil", "nintendo-64/goldeneye-007"),
        ("Hitman Blood Money", "xbox/max-payne"),
        ("Madden 97", "super-nintendo/madden-96"),
        ("NBA Live 98", "sega-genesis/nba-live-97"),
    ]
    assert [t for t, u in ok if not link_matches_title(t, url(u))] == []
    assert [t for t, u in wrong if link_matches_title(t, url(u))] == []


def test_import_distrusts_wrong_links_and_keeps_genesis(admin):
    rows = [
        {"Name": "Sonic the Hedgehog", "Sheet": "Genesis", "Link": "https://www.pricecharting.com/game/sega-genesis/sonic-the-hedgehog"},
        {"Name": "Stealth", "Sheet": "NES", "Link": "https://www.pricecharting.com/game/nes/rocket-ranger"},
        {"Name": "Ecco Tides of Time", "Sheet": "Genesis", "Link": "https://www.pricecharting.com/game/xbox/ecco-tides-of-time"},
        {"Name": "Sonic 2", "Sheet": "Genesis", "Link": "https://www.pricecharting.com/game/pal-sega-mega-drive/sonic-2"},
        {"Name": "Last Alert", "Sheet": "TG16", "Link": "https://www.pricecharting.com/game/turbografx-cd/last-alert"},
    ]  # fmt: skip
    mapping = {"title": "Name", "platform": "Sheet", "pricecharting_url": "Link"}
    res = admin.post("/api/import/validate", json={"rows": rows, "mapping": mapping}).json()["rows"]
    assert res[0]["platform"] == "Sega Genesis" and res[0]["warnings"] == []  # not Nomad
    assert "looks like a different game" in res[1]["warnings"][0]
    assert res[2]["platform"] == "Sega Genesis" and "is for Xbox, not Sega Genesis" in res[2]["warnings"][0]
    assert res[3]["platform"] == "Sega Mega Drive (PAL)"  # same console family, other region: trust the link
    assert res[4]["platform"] == "TurboGrafx-CD" and res[4]["warnings"] == []  # add-on for the same console


def _name_value_workbook() -> bytes:
    """A price-tracker layout: a tab per platform headed "Name | Value | Information | TOTAL",
    a "Consoles" tab, a summary tab, and a want list with one column per platform."""
    wb = Workbook()
    wb.active.title = "Consoles"
    tabs = {
        "Consoles": [
            ("Sega Nomad", 80, "genesis/sega-nomad"),
            ("Nintendo 64 (with expansion pack)", 63, None),
        ],
        "Genesis": [
            ("Desert Strike (Loose)", 6.95, "sega-genesis/desert-strike"),
            ("Mortal Kombat (No Manual)", 8, None),
        ],
        "N3DS": [("Pokemon Y", 23.8, None)],
        "PS2": [("Final Fantasy X (Sealed)", 11.5, None)],
        "SNES": [("R-Type, W/ Manual Loose", 12.36, None), ("Final Fantasy III W/ Guide", 242.42, None)],
    }
    for name, games in tabs.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        ws["A1"] = '=image("http://example.com/box.png")'
        for c, h in enumerate(["Name", "Value", "Information", "TOTAL", f"=SUM(B3:B{len(games) + 2})"], 1):
            ws.cell(2, c, h)
        for r, (title, value, slug) in enumerate(games, start=3):
            ws.cell(r, 1, title)
            ws.cell(r, 2, value)
            if slug:
                ws.cell(r, 3, f'=HYPERLINK("http://videogames.pricecharting.com/game/{slug}","Data Link")')
    totals = wb.create_sheet("TOTALS")
    totals.append(["CONSOLE", "COUNT", "VALUE", "COMPLETE"])
    totals.append(["NES", 17, 144.77, 0.02])
    wish = wb.create_sheet("Wishlist")
    wish.append([" NES", "Super NES"])
    wish.append(["Battletoads", "Chrono Trigger"])
    wish.append([None, "Super Metroid"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_name_value_tracker(admin):
    r = admin.post("/api/import/workbook", files={"file": ("c.xlsx", _name_value_workbook())})
    sheets = {s["name"]: s for s in r.json()["sheets"]}
    assert sheets["TOTALS"]["rows"] == []  # a summary, not items
    assert sheets["Genesis"]["headers"] == [
        "Name",
        "Value",
        "Information",
        "Information (link)",
        "Sheet name",
    ]
    assert sheets["Consoles"]["headers"] == sheets["Genesis"]["headers"]  # same layout: import together
    assert [(w["Title"], w["Platform"]) for w in sheets["Wishlist"]["rows"]] == [
        ("Battletoads", "NES"),
        ("Chrono Trigger", "Super NES"),
        ("Super Metroid", "Super NES"),
    ]

    rows = [row for n in ("Consoles", "Genesis", "N3DS", "PS2", "SNES") for row in sheets[n]["rows"]]
    headers = sheets["Genesis"]["headers"]
    mapping = admin.post("/api/import/suggest", json={"headers": headers, "rows": rows}).json()
    assert mapping == {
        "title": "Name",
        "pricecharting_url": "Information (link)",
        "platform": "Sheet name",
    }
    result = admin.post("/api/import/commit", json={"rows": rows, "mapping": mapping}).json()
    assert result["failed"] == 0 and result["created_items"] == 8
    items = {i["product"]["title"]: i for i in admin.get("/api/collection").json()["items"]}

    nomad = items["Sega Nomad"]
    assert nomad["product"]["category"] == "console" and nomad["product"]["platform"]["name"] == "Sega Nomad"
    n64 = items["Nintendo 64"]
    assert n64["product"]["platform"]["name"] == "Nintendo 64" and n64["notes"] == "With expansion pack"
    assert items["Pokemon Y"]["product"]["platform"]["name"] == "Nintendo 3DS"
    desert = items["Desert Strike"]
    assert desert["condition"] == "loose" and desert["purchase_price"] is None
    assert (
        desert["product"]["pricecharting_url"]
        == "https://www.pricecharting.com/game/sega-genesis/desert-strike"
    )
    assert not items["Mortal Kombat"]["has_manual"]
    assert items["Final Fantasy X"]["condition"] == "new"
    rtype = items["R-Type"]
    assert rtype["condition"] == "loose" and rtype["has_manual"] and not rtype["has_box"]
    assert items["Final Fantasy III"]["notes"] == "With Guide"

    wish = sheets["Wishlist"]
    mapping = admin.post(
        "/api/import/suggest", json={"headers": wish["headers"], "rows": wish["rows"]}
    ).json()
    assert mapping == {"title": "Title", "platform": "Platform", "status": "Status"}
    result = admin.post("/api/import/commit", json={"rows": wish["rows"], "mapping": mapping}).json()
    assert result["created_items"] == 3 and result["failed"] == 0
    wanted = admin.get("/api/collection", params={"status": "wishlist"}).json()["items"]
    assert {(i["product"]["title"], i["product"]["platform"]["name"]) for i in wanted} == {
        ("Battletoads", "NES"),
        ("Chrono Trigger", "Super Nintendo"),
        ("Super Metroid", "Super Nintendo"),
    }
