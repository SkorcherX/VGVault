"""Read Excel workbooks for import.

Collectors' spreadsheets rarely start at A1: they have title rows, totals above the
header, one tab per platform, and links hidden in HYPERLINK() formulas. This finds each
sheet's header row, turns cell values into plain text, and exposes hyperlinks as an extra
"<column> (link)" column so a PriceCharting link can be mapped like any other field.
"""

import re
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook

from app.services.importer import FIELDS, HEADER_HINTS, norm

SHEET_COLUMN = "Sheet name"
MAX_ROWS = 10_000
MAX_EMPTY_RUN = 200  # stop scanning a sheet after this many blank rows in a row
HEADER_SCAN_ROWS = 25
HYPERLINK_RE = re.compile(r'^=\s*HYPERLINK\(\s*"([^"]+)"', re.IGNORECASE)

# Headers that don't feed a field but still mark a header row ("Name | Value | Information").
EXTRA_HEADERS = {"value", "information", "info", "details", "description", "link", "links"}
KNOWN_HEADERS = (
    {norm(f) for f in FIELDS} | {h for hints in HEADER_HINTS.values() for h in hints} | EXTRA_HEADERS
)
TITLE_HEADERS = {"title", *HEADER_HINTS["title"]}
# Tabs laid out as one column per platform, each a list of titles (a quick want list).
LIST_SHEETS = {"wishlist", "wishlists", "wantlist", "wants", "wanted", "wishes"}


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.6f}".rstrip("0").rstrip(".")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value).strip()


def _header_row(ws) -> int | None:
    """The row (1-based) among the first few that looks most like column headers. It must name
    a title column, so summary tabs ("Console | Count | Value") aren't read as items."""
    best, best_score = None, 1
    for r, row in enumerate(ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
        names = [norm(v) for v in row if isinstance(v, str)]
        if not TITLE_HEADERS & set(names):
            continue
        score = sum(1 for n in names if n in KNOWN_HEADERS)
        if score > best_score:
            best, best_score = r, score
    return best


def _link(value_cell, formula_cell) -> str:
    if value_cell.hyperlink is not None and value_cell.hyperlink.target:
        return value_cell.hyperlink.target
    formula = formula_cell.value
    if isinstance(formula, str):
        m = HYPERLINK_RE.match(formula)
        if m:
            return m.group(1)
    return ""


def read_workbook(data: bytes) -> list[dict]:
    values_wb = load_workbook(BytesIO(data), data_only=True)  # computed values
    formulas_wb = load_workbook(BytesIO(data), data_only=False)  # to read HYPERLINK() targets
    sheets = []
    for ws in values_wb.worksheets:
        wf = formulas_wb[ws.title]
        info = {
            "name": ws.title,
            "hidden": ws.sheet_state != "visible",
            "header_row": None,
            "headers": [],
            "rows": [],
            "row_numbers": [],  # spreadsheet row of each entry in rows, for messages
        }
        sheets.append(info)
        header_row = _header_row(ws)
        if header_row is None:
            if norm(ws.title) in LIST_SHEETS:
                _read_list_sheet(ws, info)
            continue

        # Column index -> unique header name (blank headers are skipped)
        columns: dict[int, str] = {}
        seen: dict[str, int] = {}
        for c, v in enumerate(
            next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)), start=1
        ):
            name = cell_text(v)
            if not name or not isinstance(v, str):  # blank, or a total sitting in the header row
                continue
            seen[name] = seen.get(name, 0) + 1
            columns[c] = name if seen[name] == 1 else f"{name} ({seen[name]})"

        rows: list[dict[str, str]] = []
        row_numbers: list[int] = []
        link_columns: set[str] = set()
        filled: set[str] = set()
        empty_run = 0
        for r in range(header_row + 1, ws.max_row + 1):
            row = {}
            for c, name in columns.items():
                text = cell_text(ws.cell(r, c).value)
                if text:
                    row[name] = text
                    filled.add(name)
                link = _link(ws.cell(r, c), wf.cell(r, c))
                if link:
                    row[f"{name} (link)"] = link
                    link_columns.add(f"{name} (link)")
            if not row:
                empty_run += 1
                if empty_run >= MAX_EMPTY_RUN:
                    break
                continue
            empty_run = 0
            row[SHEET_COLUMN] = ws.title
            rows.append(row)
            row_numbers.append(r)
            if len(rows) >= MAX_ROWS:
                break

        info["header_row"] = header_row
        # Columns with nothing under them (a "TOTAL" label beside the header) aren't offered, so
        # tabs that differ only in such labels still import together.
        headers = [n for n in columns.values() if n in filled or f"{n} (link)" in link_columns]
        info["headers"] = [*headers, *sorted(link_columns), SHEET_COLUMN]
        info["rows"] = rows
        info["row_numbers"] = row_numbers
    return sheets


def _read_list_sheet(ws, info: dict) -> None:
    """A want list with a column per platform: the first filled row names the platforms and the
    cells below are titles. Rows come out as Title / Platform / Status like any other sheet."""
    columns: dict[int, str] = {}
    rows: list[dict[str, str]] = []
    row_numbers: list[int] = []
    for r, values in enumerate(ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS), values_only=True), start=1):
        if not columns:
            columns = {c: cell_text(v) for c, v in enumerate(values) if cell_text(v)}
            if columns:
                info["header_row"] = r
            continue
        for c, platform in columns.items():
            title = cell_text(values[c]) if c < len(values) else ""
            if title:
                rows.append(
                    {"Title": title, "Platform": platform, "Status": "wishlist", SHEET_COLUMN: ws.title}
                )
                row_numbers.append(r)
    if rows:
        info["headers"] = ["Title", "Platform", "Status", SHEET_COLUMN]
        info["rows"] = rows
        info["row_numbers"] = row_numbers
